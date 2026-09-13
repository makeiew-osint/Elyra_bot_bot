import asyncio
import base64
import html
import io
import logging
import re
import sqlite3
import tempfile
from enum import Enum
from pathlib import Path

from aiogram import Bot, Dispatcher, F, Router
from aiogram.exceptions import TelegramAPIError, TelegramBadRequest
from aiogram.filters import Command, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    CallbackQuery,
    BotCommand,
    BufferedInputFile,
    FSInputFile,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    Message,
)
from huggingface_hub import InferenceClient
from google import genai
from google.genai import types

from config import (
    EDIT_MODEL,
    GEMINI_MODEL,
    IMAGE_FAST_MODEL,
    IMAGE_REALISTIC_MODEL,
    IMAGE_MODEL,
    OCR_MODEL,
    REASONING_MODEL,
    SUPPORT_USERNAME,
    TEXT_MODEL,
    Settings,
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
router = Router()
HISTORY_DB = Path(__file__).with_name("elyra_history.sqlite3")
CONTEXT_MESSAGES = 24


class Mode(str, Enum):
    CHAT = "chat"
    CODE = "code"
    AGENT = "agent"
    OCR = "ocr"
    IMAGE = "image"
    EDIT = "edit"
    SEARCH = "search"


class UserFlow(StatesGroup):
    waiting_for_prompt = State()
    waiting_for_image_prompt = State()
    waiting_for_edit_image = State()
    waiting_for_edit_prompt = State()


def main_menu() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="🧠 Умный чат", callback_data="mode:chat"),
                InlineKeyboardButton(text="💻 Решить код", callback_data="mode:code"),
            ],
            [
                InlineKeyboardButton(text="🤖 Агент", callback_data="mode:agent"),
                InlineKeyboardButton(text="📷 Решить фото / 📄 файл", callback_data="mode:ocr"),
            ],
            [InlineKeyboardButton(text="🌐 Поиск в интернете", callback_data="mode:search")],
            [
                InlineKeyboardButton(text="🎨 Создать картинку", callback_data="mode:image"),
                InlineKeyboardButton(text="✏️ Изменить фото", callback_data="mode:edit"),
            ],
            [
                InlineKeyboardButton(text="🗑 Очистить историю", callback_data="history:clear"),
            ],
            [
                InlineKeyboardButton(text="ℹ️ О боте", callback_data="about"),
                InlineKeyboardButton(text="🆘 Поддержка", url=f"https://t.me/{SUPPORT_USERNAME}"),
            ],
        ]
    )


def back_menu() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="✖️ Отмена", callback_data="cancel"),
                InlineKeyboardButton(text="⌂ Главное меню", callback_data="menu"),
            ],
            [InlineKeyboardButton(text="🆘 Поддержка", url=f"https://t.me/{SUPPORT_USERNAME}")],
        ]
    )


ABOUT_TEXT = (
    "✨ <b>Elyra</b> — твой AI-помощник в Telegram.\n\n"
    "🧠 <b>Умный чат</b> — вопросы и объяснения\n"
    "💻 <b>Решить код</b> — программирование\n"
    "🤖 <b>Агент</b> — планы и сложные задачи\n"
    "📷 <b>Решить фото</b> — условие с картинки + решение\n"
    "🌐 <b>Поиск в интернете</b> — актуальные сведения через Gemini\n"
    "🎨 <b>Создать картинку</b> — генерация\n"
    "✏️ <b>Изменить фото</b> — редактирование\n\n"
    "Выберите нужный режим — я помогу разобраться быстро и понятно.\n"
    "🆘 Поддержка: @Makeiew"
)


def prompt_for_mode(mode: Mode) -> str:
    return {
        Mode.CHAT: "Напишите вопрос или задачу:",
        Mode.CODE: "Опишите задачу по коду. Укажите язык и ожидаемый результат:",
        Mode.AGENT: "Опишите цель. Я разложу ее на шаги и предложу решение:",
        Mode.OCR: "Отправьте фото или документ для распознавания текста:",
        Mode.SEARCH: "Напишите, что найти в интернете:",
        Mode.IMAGE: "Опишите изображение, которое нужно создать:",
        Mode.EDIT: "Сначала отправьте изображение, затем напишите, что изменить:",
    }[mode]


def init_history() -> None:
    with sqlite3.connect(HISTORY_DB) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS conversation_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        connection.commit()


def load_history(user_id: int) -> list[tuple[str, str]]:
    with sqlite3.connect(HISTORY_DB) as connection:
        rows = connection.execute(
            """
            SELECT role, content FROM conversation_history
            WHERE user_id = ? ORDER BY id DESC LIMIT ?
            """,
            (user_id, CONTEXT_MESSAGES),
        ).fetchall()
    return list(reversed(rows))


def save_history(user_id: int, role: str, content: str) -> None:
    with sqlite3.connect(HISTORY_DB) as connection:
        connection.execute(
            "INSERT INTO conversation_history (user_id, role, content) VALUES (?, ?, ?)",
            (user_id, role, content),
        )
        connection.commit()


def clear_history(user_id: int) -> None:
    with sqlite3.connect(HISTORY_DB) as connection:
        connection.execute("DELETE FROM conversation_history WHERE user_id = ?", (user_id,))
        connection.commit()


def history_prompt(history: list[tuple[str, str]], prompt: str) -> str:
    if not history:
        return prompt
    transcript = "\n".join(f"{role}: {content}" for role, content in history)
    return f"Контекст предыдущего диалога:\n{transcript}\n\nНовый запрос пользователя:\n{prompt}"


async def hf_text(
    client: InferenceClient,
    prompt: str,
    mode: Mode,
    history: list[tuple[str, str]] | None = None,
) -> str:
    system = {
        Mode.CHAT: "Отвечай на русском ясно и полезно.",
        Mode.CODE: "Ты опытный разработчик. Дай рабочий код и кратко объясни решение.",
        Mode.AGENT: "Ты агент-планировщик. Разбей задачу на безопасные проверяемые шаги.",
        Mode.SEARCH: "Отвечай по-русски, используя актуальные сведения из интернета.",
    }[mode]
    result = await asyncio.to_thread(
        client.chat_completion,
        messages=[
            {"role": "system", "content": system},
            {
                "role": "user",
                "content": (
                    "Оформи ответ аккуратно: используй Markdown-заголовки, списки, "
                    "таблицы и блоки кода. Формулы пиши без $...$.\n\n"
                    + history_prompt(history or [], prompt)
                ),
            },
        ],
        model=TEXT_MODEL if mode not in (Mode.AGENT, Mode.SEARCH) else REASONING_MODEL,
        max_tokens=4096,
        temperature=0.7,
    )
    return result.choices[0].message.content


def choose_text_backend(settings: Settings, mode: Mode, prompt: str) -> str:
    if mode == Mode.CODE:
        return "deepseek"
    if mode == Mode.AGENT:
        return "glm"
    if settings.gemini_api_key and len(prompt) > 1200:
        return "gemini"
    return "gemini" if settings.gemini_api_key else "deepseek"


async def gemini_text(
    settings: Settings,
    prompt: str,
    mode: Mode,
    history: list[tuple[str, str]] | None = None,
) -> str:
    client = genai.Client(api_key=settings.gemini_api_key)
    instruction = {
        Mode.CHAT: "Отвечай по-русски ясно, точно и полезно.",
        Mode.CODE: "Реши задачу по программированию. Дай рабочий код и объяснение.",
        Mode.AGENT: "Разбей задачу на проверяемые шаги и предложи надежное решение.",
        Mode.SEARCH: "Найди актуальную информацию в интернете и укажи источники.",
    }[mode]
    result = await asyncio.to_thread(
        client.models.generate_content,
        model=GEMINI_MODEL,
        contents=(
            f"{instruction}\n"
            "Используй Markdown: заголовки, списки, жирный текст, блоки кода и таблицы. "
            "Формулы пиши читабельно и не используй LaTeX-делимитеры $...$.\n\n"
            f"{history_prompt(history or [], prompt)}"
        ),
    )
    return result.text


async def gemini_search(settings: Settings, prompt: str) -> str:
    """Use Gemini grounding when the installed SDK/model supports Google Search."""
    client = genai.Client(api_key=settings.gemini_api_key)
    result = await asyncio.to_thread(
        client.models.generate_content,
        model=GEMINI_MODEL,
        contents=f"Ответь по-русски на запрос и укажи ссылки на источники:\n{prompt}",
        config=types.GenerateContentConfig(
            tools=[types.Tool(google_search=types.GoogleSearch())]
        ),
    )
    return result.text


async def gemini_transcribe(settings: Settings, content: bytes, mime_type: str) -> str:
    client = genai.Client(api_key=settings.gemini_api_key)
    result = await asyncio.to_thread(
        client.models.generate_content,
        model=GEMINI_MODEL,
        contents=[
            types.Part.from_bytes(data=content, mime_type=mime_type),
            "Точно расшифруй голосовое сообщение на русском языке. Верни только текст.",
        ],
    )
    return (result.text or "").strip()


async def hf_transcribe(client: InferenceClient, content: bytes) -> str:
    result = await asyncio.to_thread(
        client.automatic_speech_recognition,
        content,
        model="openai/whisper-large-v3-turbo",
    )
    return (getattr(result, "text", "") or "").strip()


async def hf_image(client: InferenceClient, prompt: str, model: str) -> bytes:
    image = await asyncio.to_thread(client.text_to_image, prompt=prompt, model=model)
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
        image.save(handle, format="PNG")
        return Path(handle.name).read_bytes()


def choose_image_models(prompt: str) -> list[str]:
    lowered = prompt.lower()
    realistic_words = (
        "фотореалист", "реалистич", "портрет", "фото", "photorealistic",
        "realistic", "portrait", "product", "товар", "архитектур",
    )
    fast_words = (
        "быстро", "просто", "икон", "логотип", "стикер", "мультяш",
        "аниме", "скетч", "эскиз", "simple", "icon", "logo", "sticker",
        "anime", "cartoon",
    )
    if any(word in lowered for word in realistic_words):
        preferred = IMAGE_REALISTIC_MODEL
    elif any(word in lowered for word in fast_words):
        preferred = IMAGE_FAST_MODEL
    else:
        preferred = IMAGE_MODEL
    return [preferred] + [
        model for model in (IMAGE_MODEL, IMAGE_REALISTIC_MODEL, IMAGE_FAST_MODEL)
        if model != preferred
    ]


async def generate_image_with_fallback(
    client: InferenceClient, prompt: str
) -> tuple[bytes, str]:
    last_error: Exception | None = None
    for model in choose_image_models(prompt):
        try:
            return await hf_image(client, prompt, model), model
        except Exception as error:
            last_error = error
            logging.warning("Image model %s failed; trying next model", model, exc_info=True)
    raise RuntimeError("All image generation models failed") from last_error


def temporary_image_path(content: bytes) -> Path:
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
        handle.write(content)
        return Path(handle.name)


def data_uri_for_image(content: bytes) -> str:
    if content.startswith(b"\x89PNG"):
        mime_type = "image/png"
    elif content.startswith(b"GIF8"):
        mime_type = "image/gif"
    elif content.startswith(b"RIFF") and content[8:12] == b"WEBP":
        mime_type = "image/webp"
    else:
        mime_type = "image/jpeg"
    return f"data:{mime_type};base64," + base64.b64encode(content).decode("ascii")


async def hf_edit(client: InferenceClient, content: bytes, prompt: str) -> bytes:
    image = await asyncio.to_thread(
        client.image_to_image,
        image=content,
        prompt=prompt,
        model=EDIT_MODEL,
    )
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
        image.save(handle, format="PNG")
        return Path(handle.name).read_bytes()


async def hf_ocr(client: InferenceClient, content: bytes) -> str:
    result = await asyncio.to_thread(
        client.chat_completion,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": (
                            "Распознай весь текст на изображении без пропусков. "
                            "Сохрани номера, формулы и переносы строк."
                        ),
                    },
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": data_uri_for_image(content)
                        },
                    },
                ],
            }
        ],
        model=OCR_MODEL,
        max_tokens=2048,
    )
    return result.choices[0].message.content


async def hf_vision_answer(
    client: InferenceClient, content: bytes, prompt: str
) -> str:
    result = await asyncio.to_thread(
        client.chat_completion,
        messages=[
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {
                        "type": "image_url",
                        "image_url": {"url": data_uri_for_image(content)},
                    },
                ],
            }
        ],
        model=TEXT_MODEL,
        max_tokens=2048,
    )
    return result.choices[0].message.content


async def gemini_vision_answer(
    settings: Settings, content: bytes, prompt: str
) -> str:
    client = genai.Client(api_key=settings.gemini_api_key)
    result = await asyncio.to_thread(
        client.models.generate_content,
        model=GEMINI_MODEL,
        contents=[
            types.Part.from_bytes(
                data=content,
                mime_type=data_uri_for_image(content).split(";", 1)[0][5:],
            ),
            prompt,
        ],
    )
    return result.text


async def gemini_document_answer(
    settings: Settings, content: bytes, mime_type: str, prompt: str
) -> str:
    client = genai.Client(api_key=settings.gemini_api_key)
    result = await asyncio.to_thread(
        client.models.generate_content,
        model=GEMINI_MODEL,
        contents=[types.Part.from_bytes(data=content, mime_type=mime_type), prompt],
    )
    return (result.text or "").strip()


async def extract_document_text(content: bytes, name: str, mime: str) -> str:
    if name.endswith(".txt") or mime.startswith("text/"):
        return content.decode("utf-8", errors="replace")[:120000]
    if name.endswith(".docx"):
        from docx import Document
        document = Document(io.BytesIO(content))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts.append("\n".join(
                "\t".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            ))
        return "\n".join(parts)[:120000]
    if name.endswith((".xlsx", ".xlsm")) or "spreadsheet" in mime:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = []
        for sheet in workbook.worksheets:
            rows.append(f"[Лист: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                rows.append("\t".join("" if value is None else str(value) for value in row))
        return "\n".join(rows)[:120000]
    raise RuntimeError("Поддерживаются PDF, DOCX, TXT и XLSX.")


async def send_error(message: Message, error: Exception) -> None:
    logging.exception("Hugging Face request failed", exc_info=error)
    error_text = str(error)
    if "402" in error_text or "depleted your monthly included credits" in error_text:
        text = (
            "💳 Лимит генерации изображений Hugging Face исчерпан.\n\n"
            "Пополните кредиты или подключите тариф Hugging Face, "
            "после чего генерация снова заработает.\n"
            "Текстовый чат и решение задач продолжают работать."
        )
    else:
        text = (
            "Не удалось обработать запрос. Проверьте доступность сервиса "
            "или обратитесь в поддержку @Makeiew."
        )
    await message.answer(
        text,
        reply_markup=back_menu(),
    )


def _format_table(lines: list[str]) -> str:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and not all(set(cell) <= {"-", ":", " "} for cell in cells):
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    sizes = [max(len(row[index]) for row in rows) for index in range(width)]
    return "\n".join("  ".join(cell.ljust(sizes[index]) for index, cell in enumerate(row)) for row in rows)


def format_answer(text: str) -> str:
    text = text.replace("\\(", "").replace("\\)", "").replace("\\[", "").replace("\\]", "")
    output: list[str] = []
    in_code = False
    code_lines: list[str] = []
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.strip().startswith("```"):
            if in_code:
                output.append(f"<pre>{html.escape(chr(10).join(code_lines))}</pre>")
                code_lines = []
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if "|" in line and index + 1 < len(lines) and "|" in lines[index + 1]:
            table_lines = []
            while index < len(lines) and "|" in lines[index]:
                table_lines.append(lines[index])
                index += 1
            table = _format_table(table_lines)
            if table:
                output.append(f"<pre>{html.escape(table)}</pre>")
                continue
        escaped = html.escape(line)
        escaped = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)
        escaped = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<i>\1</i>", escaped)
        escaped = re.sub(r"`([^`]+)`", r"<code>\1</code>", escaped)
        escaped = re.sub(r"^#{1,6}\s+(.+)$", r"<b>\1</b>", escaped)
        output.append(escaped)
        index += 1
    if in_code:
        output.append(f"<pre>{html.escape(chr(10).join(code_lines))}</pre>")
    result = "\n".join(output).strip()
    return result or "Пустой ответ."


async def send_answer(message: Message, answer: str) -> None:
    clean_answer = answer.strip()
    formatted = format_answer(clean_answer)
    if len(formatted) > 4000:
        await message.answer(clean_answer[:3900], reply_markup=back_menu())
    else:
        try:
            await message.answer(formatted, reply_markup=back_menu(), parse_mode="HTML")
        except TelegramBadRequest:
            logging.warning("Telegram rejected formatted answer; sending plain text", exc_info=True)
            await message.answer(clean_answer[:4000], reply_markup=back_menu())
    document = BufferedInputFile(clean_answer.encode("utf-8"), filename="elyra_answer.txt")
    await message.answer_document(document, caption="📄 Полный ответ в TXT")


async def thinking(message: Message) -> Message:
    return await message.answer("⏳ Думаю над ответом…")


async def clear_thinking(status: Message) -> None:
    try:
        await status.delete()
    except TelegramAPIError:
        logging.debug("Could not remove thinking status", exc_info=True)


async def send_main_menu(message: Message, caption: str) -> None:
    image = Path(__file__).with_name("assets").joinpath("elyra_welcome.png")
    if image.exists():
        await message.answer_photo(
            FSInputFile(image),
            caption=caption,
            reply_markup=main_menu(),
            parse_mode="HTML",
        )
    else:
        await message.answer(caption, reply_markup=main_menu(), parse_mode="HTML")


@router.message(CommandStart())
async def start(message: Message, state: FSMContext) -> None:
    await state.clear()
    await send_main_menu(
        message,
        "✨ <b>Добро пожаловать в Elyra!</b>\n\n"
        "Я помогу разобраться с задачей, объяснить сложную тему, "
        "написать код или решить пример с фотографии.\n\n"
        "💬 Напиши вопрос или отправь фото — я сразу начну помогать.\n\n"
        "👇 Выбери нужный режим:",
    )


@router.message(Command("help"))
async def help_command(message: Message) -> None:
    await message.answer(
        "Выберите режим в меню и отправьте запрос. Можно отправить голос, фото, PDF, DOCX, TXT или XLSX.\n"
        "Команда /cancel сбрасывает текущий режим.",
        reply_markup=main_menu(),
    )


@router.message(Command("about"))
async def about_command(message: Message) -> None:
    await message.answer(ABOUT_TEXT, reply_markup=main_menu(), parse_mode="HTML")


@router.message(Command("cancel"))
async def cancel_command(message: Message, state: FSMContext) -> None:
    await state.clear()
    await send_main_menu(
        message,
        "✅ Текущий запрос отменён.\n\nВыберите, что сделаем дальше:",
    )


@router.message(Command("clearhistory"))
async def clear_history_command(message: Message, state: FSMContext) -> None:
    await asyncio.to_thread(clear_history, message.from_user.id)
    await state.clear()
    await send_main_menu(
        message,
        "🗑 История диалога удалена.\n\nВыберите режим для нового разговора:",
    )


@router.callback_query(F.data == "menu")
async def menu_callback(callback: CallbackQuery, state: FSMContext) -> None:
    await state.clear()
    await send_main_menu(
        callback.message,
        "✨ <b>Elyra</b>\n\nВыберите режим работы:",
    )
    await callback.answer()


@router.callback_query(F.data == "cancel")
async def cancel_callback(callback: CallbackQuery, state: FSMContext) -> None:
    await state.clear()
    await send_main_menu(
        callback.message,
        "✅ Текущий запрос отменён.\n\nВыберите, что сделаем дальше:",
    )
    await callback.answer("Отменено")


@router.callback_query(F.data == "history:clear")
async def clear_history_callback(callback: CallbackQuery, state: FSMContext) -> None:
    clear_history(callback.from_user.id)
    await state.clear()
    await send_main_menu(
        callback.message,
        "🗑 <b>История очищена</b>\n\n"
        "Начнём новый диалог. Выберите нужный режим:",
    )
    await callback.answer("История удалена")


@router.callback_query(F.data == "about")
async def about_callback(callback: CallbackQuery) -> None:
    await callback.message.answer(ABOUT_TEXT, reply_markup=back_menu(), parse_mode="HTML")
    await callback.answer()


@router.callback_query(F.data.startswith("mode:"))
async def mode_callback(callback: CallbackQuery, state: FSMContext) -> None:
    mode = Mode(callback.data.split(":", 1)[1])
    await state.clear()
    if mode in (Mode.IMAGE, Mode.EDIT):
        await callback.message.answer(
            "🔧 <b>Технические работы</b>\n\n"
            "Функция изображений временно недоступна. "
            "Остальные возможности Elyra работают.",
            reply_markup=back_menu(),
            parse_mode="HTML",
        )
        await callback.answer("Технические работы")
        return
    if mode in (Mode.IMAGE,):
        await state.set_state(UserFlow.waiting_for_image_prompt)
    elif mode == Mode.OCR:
        await state.set_state(UserFlow.waiting_for_prompt)
    else:
        await state.set_state(UserFlow.waiting_for_prompt)
    await state.update_data(mode=mode.value)
    await callback.message.answer(prompt_for_mode(mode), reply_markup=back_menu())
    await callback.answer()


@router.message(UserFlow.waiting_for_prompt, F.text)
async def text_request(message: Message, state: FSMContext, settings: Settings) -> None:
    data = await state.get_data()
    mode = Mode(data["mode"])
    if mode == Mode.OCR:
        await message.answer("Для OCR отправьте изображение или PDF.", reply_markup=back_menu())
        return
    status = await thinking(message)
    try:
        history = await asyncio.to_thread(load_history, message.from_user.id)
        if mode == Mode.SEARCH:
            if not settings.gemini_api_key:
                raise RuntimeError("Для поиска в интернете требуется GEMINI_API_KEY.")
            answer = await gemini_search(settings, message.text)
        elif choose_text_backend(settings, mode, message.text) == "gemini":
            try:
                answer = await gemini_text(settings, message.text, mode, history)
            except Exception:
                logging.warning("Gemini failed; falling back to DeepSeek", exc_info=True)
                answer = await hf_text(
                    InferenceClient(token=settings.hf_token), message.text, mode, history
                )
        else:
            answer = await hf_text(
                InferenceClient(token=settings.hf_token), message.text, mode, history
            )
        await asyncio.to_thread(save_history, message.from_user.id, "user", message.text)
        await asyncio.to_thread(save_history, message.from_user.id, "assistant", answer)
        await send_answer(message, answer)
    except Exception as error:
        await send_error(message, error)
    finally:
        await clear_thinking(status)


@router.message(F.voice)
async def voice_request(message: Message, bot: Bot, settings: Settings) -> None:
    if not settings.gemini_api_key:
        await message.answer(
            "🎙 Для расшифровки голосовых сообщений требуется GEMINI_API_KEY.",
            reply_markup=back_menu(),
        )
        return
    status = await thinking(message)
    try:
        file = await bot.get_file(message.voice.file_id)
        buffer = await bot.download_file(file.file_path)
        content = buffer.read()
        try:
            prompt = await asyncio.wait_for(
                gemini_transcribe(settings, content, "audio/ogg"),
                timeout=90,
            )
        except Exception:
            logging.warning("Gemini transcription failed; using Whisper fallback", exc_info=True)
            prompt = await asyncio.wait_for(
                hf_transcribe(InferenceClient(token=settings.hf_token), content),
                timeout=90,
            )
        if not prompt:
            raise RuntimeError("Расшифровка голосового сообщения пуста.")
        history = await asyncio.to_thread(load_history, message.from_user.id)
        try:
            answer = await asyncio.wait_for(
                gemini_text(settings, prompt, Mode.CHAT, history),
                timeout=90,
            )
        except Exception:
            logging.warning("Gemini voice answer failed; falling back to DeepSeek", exc_info=True)
            answer = await asyncio.wait_for(
                hf_text(
                    InferenceClient(token=settings.hf_token),
                    prompt,
                    Mode.CHAT,
                    history,
                ),
                timeout=90,
            )
        await asyncio.to_thread(save_history, message.from_user.id, "user", "[Голос]\n" + prompt)
        await asyncio.to_thread(save_history, message.from_user.id, "assistant", answer)
        await message.answer(f"📝 <b>Расшифровка:</b>\n{html.escape(prompt)}", parse_mode="HTML")
        await send_answer(message, answer)
    except Exception as error:
        await send_error(message, error)
    finally:
        await clear_thinking(status)


@router.message(UserFlow.waiting_for_prompt, F.photo)
async def ocr_photo(message: Message, bot: Bot, state: FSMContext, settings: Settings) -> None:
    data = await state.get_data()
    status = await thinking(message)
    try:
        mode = Mode(data["mode"])
        file = await bot.get_file(message.photo[-1].file_id)
        buffer = await bot.download_file(file.file_path)
        content = buffer.read()
        client = InferenceClient(token=settings.hf_token)
        if settings.gemini_api_key:
            try:
                extracted = await gemini_vision_answer(
                    settings,
                    content,
                    "Распознай текст и условие задачи на фотографии. "
                    "Сохрани формулы, номера и все важные детали.",
                )
            except Exception:
                logging.warning("Gemini vision failed; trying Hugging Face OCR", exc_info=True)
                try:
                    extracted = await hf_ocr(client, content)
                except Exception:
                    logging.warning(
                        "GLM-OCR failed; using multimodal DeepSeek fallback", exc_info=True
                    )
                    extracted = await hf_vision_answer(
                        client,
                        content,
                        "Точно распознай текст и условие задачи на фотографии. "
                        "Сохрани формулы, номера и все важные детали.",
                    )
        else:
            try:
                extracted = await hf_ocr(client, content)
            except Exception:
                logging.warning("GLM-OCR failed; using multimodal DeepSeek fallback", exc_info=True)
                extracted = await hf_vision_answer(
                    client,
                    content,
                    "Точно распознай текст на фотографии. Сохрани условие задачи, "
                    "формулы, номера и все важные детали.",
                )
        if not extracted.strip():
            await message.answer("Не удалось распознать текст на фото.", reply_markup=back_menu())
            return
        if mode == Mode.OCR:
            answer = extracted
        else:
            prompt = (
                f"Реши задачу, распознанную с фотографии. "
                f"Покажи ход решения и итоговый ответ.\n\n{extracted}"
            )
            if settings.gemini_api_key:
                try:
                    history = await asyncio.to_thread(load_history, message.from_user.id)
                    answer = await gemini_text(settings, prompt, mode, history)
                except Exception:
                    logging.warning("Gemini failed; falling back to DeepSeek", exc_info=True)
                    answer = await hf_text(client, prompt, mode, history)
            else:
                history = await asyncio.to_thread(load_history, message.from_user.id)
                answer = await hf_text(client, prompt, mode, history)
        await asyncio.to_thread(save_history, message.from_user.id, "user", "[Изображение]\n" + extracted)
        await asyncio.to_thread(save_history, message.from_user.id, "assistant", answer)
        await send_answer(message, answer)
    except Exception as error:
        await send_error(message, error)
    finally:
        await clear_thinking(status)


@router.message(UserFlow.waiting_for_image_prompt, F.photo)
async def image_question_photo(
    message: Message, bot: Bot, state: FSMContext, settings: Settings
) -> None:
    """Treat a photo with a caption as a vision question, not image generation."""
    question = (message.caption or "").strip()
    if not question:
        await message.answer(
            "Напишите вопрос подписью к фотографии, например: "
            "«Что это и сколько примерно стоит?»",
            reply_markup=back_menu(),
        )
        return

    status = await thinking(message)
    try:
        file = await bot.get_file(message.photo[-1].file_id)
        buffer = await bot.download_file(file.file_path)
        content = buffer.read()
        prompt = (
            "Ответь на вопрос пользователя по фотографии. "
            "Внимательно опиши, что видно на изображении. Если спрашивают цену, "
            "не выдумывай точную стоимость: укажи ориентировочный диапазон, "
            "от чего он зависит, и какие данные нужны для точной оценки. "
            "Если просят ссылку на магазин, честно скажи, что не можешь проверить "
            "актуальные цены без поиска в интернете.\n\n"
            f"Вопрос пользователя: {question}"
        )
        if settings.gemini_api_key:
            try:
                answer = await gemini_vision_answer(settings, content, prompt)
            except Exception:
                logging.warning("Gemini vision failed; using DeepSeek vision", exc_info=True)
                answer = await hf_vision_answer(
                    InferenceClient(token=settings.hf_token), content, prompt
                )
        else:
            answer = await hf_vision_answer(
                InferenceClient(token=settings.hf_token), content, prompt
            )
        await asyncio.to_thread(
            save_history,
            message.from_user.id,
            "user",
            f"[Вопрос по изображению]\n{question}",
        )
        await asyncio.to_thread(save_history, message.from_user.id, "assistant", answer)
        await send_answer(message, answer)
        await state.set_state(UserFlow.waiting_for_image_prompt)
    except Exception as error:
        await send_error(message, error)
    finally:
        await clear_thinking(status)


@router.message(UserFlow.waiting_for_prompt, F.document)
async def ocr_document(message: Message, bot: Bot, state: FSMContext, settings: Settings) -> None:
    data = await state.get_data()
    if data.get("mode") != Mode.OCR.value:
        await message.answer("В этом режиме нужен текстовый запрос.", reply_markup=back_menu())
        return
    status = await thinking(message)
    try:
        file = await bot.get_file(message.document.file_id)
        buffer = await bot.download_file(file.file_path)
        content = buffer.read()
        name = (message.document.file_name or "").lower()
        mime = message.document.mime_type or ""
        if mime == "application/pdf" or name.endswith(".pdf"):
            if not settings.gemini_api_key:
                raise RuntimeError("Анализ PDF требует GEMINI_API_KEY.")
            answer = await gemini_document_answer(
                settings, content, "application/pdf",
                "Извлеки текст из PDF и кратко объясни его содержание.",
            )
        else:
            extracted = await extract_document_text(content, name, mime)
            if not extracted.strip():
                raise RuntimeError("В документе не найден текст.")
            history = await asyncio.to_thread(load_history, message.from_user.id)
            answer = await gemini_text(
                settings, "Проанализируй документ:\n\n" + extracted,
                Mode.CHAT, history
            ) if settings.gemini_api_key else await hf_text(
                InferenceClient(token=settings.hf_token),
                "Проанализируй документ:\n\n" + extracted, Mode.CHAT, history
            )
        await send_answer(message, answer or "Текст не найден.")
    except Exception as error:
        await send_error(message, error)
    finally:
        await clear_thinking(status)


@router.message(UserFlow.waiting_for_image_prompt, F.text)
async def image_request(message: Message, state: FSMContext, settings: Settings) -> None:
    status = await thinking(message)
    try:
        client = InferenceClient(token=settings.hf_token)
        history = await asyncio.to_thread(load_history, message.from_user.id)
        previous_prompts = [
            content.removeprefix("[Генерация изображения]\n")
            for role, content in history
            if role == "user" and content.startswith("[Генерация изображения]")
        ]
        prompt = message.text
        if previous_prompts:
            prompt = (
                "Это продолжение работы над изображением. Учти предыдущее описание:\n"
                + previous_prompts[-1]
                + "\n\nНовое изменение пользователя:\n"
                + message.text
            )
        image, _model = await generate_image_with_fallback(client, prompt)
        path = temporary_image_path(image)
        await message.answer_photo(
            FSInputFile(path),
            caption="Готово ✨",
            reply_markup=back_menu(),
        )
        await asyncio.to_thread(
            save_history, message.from_user.id, "user", "[Генерация изображения]\n" + message.text
        )
        await asyncio.to_thread(
            save_history, message.from_user.id, "assistant", "[Изображение создано]"
        )
        path.unlink(missing_ok=True)
        await state.set_state(UserFlow.waiting_for_image_prompt)
    except Exception as error:
        await send_error(message, error)
    finally:
        await clear_thinking(status)


@router.message(UserFlow.waiting_for_edit_image, F.photo)
async def edit_image_received(message: Message, state: FSMContext) -> None:
    await state.update_data(file_id=message.photo[-1].file_id)
    await state.set_state(UserFlow.waiting_for_edit_prompt)
    await message.answer("Теперь напишите, что изменить на изображении:", reply_markup=back_menu())


@router.message(UserFlow.waiting_for_edit_image)
async def edit_image_expected(message: Message) -> None:
    await message.answer("Отправьте изображение как фото.", reply_markup=back_menu())


@router.message(UserFlow.waiting_for_edit_prompt, F.text)
async def edit_request(
    message: Message, bot: Bot, state: FSMContext, settings: Settings
) -> None:
    await message.answer(
        "Запрос на редактирование принят, обрабатываю изображение…"
    )
    try:
        data = await state.get_data()
        file = await bot.get_file(data["file_id"])
        buffer = await bot.download_file(file.file_path)
        client = InferenceClient(token=settings.hf_token)
        image = await hf_edit(client, buffer.read(), message.text)
        path = temporary_image_path(image)
        await message.answer_photo(FSInputFile(path), caption="Готово ✨", reply_markup=back_menu())
        path.unlink(missing_ok=True)
        await state.clear()
    except Exception as error:
        await send_error(message, error)


async def main() -> None:
    settings = Settings.from_env()
    init_history()
    bot = Bot(settings.telegram_token)
    try:
        await bot.set_my_commands(
            [
                BotCommand(command="start", description="Главное меню"),
                BotCommand(command="help", description="Помощь"),
                BotCommand(command="about", description="О боте"),
                BotCommand(command="cancel", description="Отменить режим"),
                BotCommand(command="clearhistory", description="Очистить историю"),
            ]
        )
    except TelegramAPIError:
        logging.warning("Bot command setup skipped; Telegram rate limit or API error")
    dp = Dispatcher()
    dp["settings"] = settings
    dp.include_router(router)
    await bot.delete_webhook(drop_pending_updates=True)
    await dp.start_polling(bot)


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except (KeyboardInterrupt, SystemExit):
        logging.info("Bot stopped")
